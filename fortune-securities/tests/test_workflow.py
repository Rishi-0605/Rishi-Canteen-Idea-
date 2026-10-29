"""End-to-end workflow on the SYNTHETIC sample data:
import holdings -> validate/reconcile -> analyse a company -> cited report ->
save record -> re-analyse after new disclosures -> compare.
"""

import json
import re
import sqlite3
import unittest

from helpers import Workspace
from openpyxl import load_workbook

from fortune_research.agent import ResearchAgent, ResolutionError


class Workflow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ws = Workspace()
        cls.s1 = cls.ws.settings(1)
        cls.agent1 = ResearchAgent(cls.s1, as_of=Workspace.as_of(1))
        cls.port = cls.agent1.analyse_portfolio()
        cls.r1 = cls.agent1.analyse_company("Sample Manufacturing")
        cls.s2 = cls.ws.settings(2)
        cls.agent2 = ResearchAgent(cls.s2, as_of=Workspace.as_of(2))
        cls.changed = cls.agent2.what_changed("SAMPLEMFG")
        cls.r2 = cls.changed["record"]

    @classmethod
    def tearDownClass(cls):
        cls.ws.cleanup()

    # 1-2: holdings import, validation and reconciliation
    def test_portfolio_import_and_reconciliation(self):
        v = self.port["view"]
        self.assertEqual(len(v["positions"]), 7)
        self.assertEqual(sorted(u["symbol"] for u in v["unpriced"]), ["SAMPLENODATA", "UNKNOWNXYZ"])
        msgs = " | ".join(i["message"] for i in self.port["issues"])
        self.assertIn("exactly 2x", msgs)                 # bonus not in transactions
        self.assertIn("trade id T1003 repeated", msgs)     # duplicate row caught
        self.assertIn("SAMPLEPHARMA: average cost", msgs)  # 3% cost difference
        self.assertEqual(self.port["reconciliation"]["matched"], 5)
        self.assertAlmostEqual(sum(r["weight"] for r in v["positions"]), 1.0)
        self.assertTrue(all("****" in a for a in v["accounts"]))  # account ids masked
        self.assertIsNotNone(v["total"]["xirr"])
        self.assertIsNotNone(v["total"]["benchmark_xirr"])
        self.assertTrue(self.port["synthetic"])

    # 3: analysis numbers are traceable to the input files
    def test_company_numbers(self):
        m = self.r1["metrics"]
        self.assertEqual(self.r1["company"]["isin"], "INEZ9M101010")
        self.assertAlmostEqual(m["ttm.eps"]["value"], 24.1)                   # 5.6 + 6.0 + 6.5 + 6.0
        self.assertAlmostEqual(m["val.price"]["value"], 690.0)
        self.assertAlmostEqual(m["val.pe"]["value"], 690 / 24.1)
        self.assertAlmostEqual(m["val.market_cap"]["value"], 6900.0)          # 690 x 10 cr shares
        self.assertAlmostEqual(m["fy.roce"]["value"], 329 / ((1210 + 170 + 1047 + 190) / 2) * 100)
        self.assertAlmostEqual(m["growth.revenue_cagr_3y"]["value"], ((1950 / 1380) ** (1 / 3) - 1) * 100)
        self.assertEqual(self.r1["data_issues"], [])  # sample statements reconcile
        self.assertEqual(self.r1["summary"]["fundamentals"], "strong")

    def test_disclosures_after_as_of_are_excluded(self):
        # stage-1 analysis must not see the Q2FY27 results published 2026-10-27
        self.assertNotIn("Q2FY27", self.r1["periods"])
        self.assertEqual(self.r1["data_as_of"]["prices"], "2026-10-20")

    # 4: report is cited and labelled
    def test_report_citations(self):
        ids = {s["id"] for s in self.r1["sources"]}
        for k, mt in self.r1["metrics"].items():
            for sid in mt["source_ids"]:
                self.assertIn(sid, ids, k)
        for f in self.r1["findings"]:
            for sid in f["source_ids"]:
                self.assertIn(sid, ids, f["code"])
        with open(self.r1["report_md_path"], encoding="utf-8") as fh:
            md = fh.read()
        self.assertIn("SYNTHETIC SAMPLE DATA", md)
        n_sources = len(self.r1["sources"])
        cited = {int(x) for x in re.findall(r"\[(\d+)\]", md)}
        self.assertTrue(cited and max(cited) <= n_sources)
        self.assertIn("## 13. Sources", md)
        self.assertIn("not investment advice", md)
        wb = load_workbook(self.r1["report_xlsx_path"])
        self.assertIn("Sources", wb.sheetnames)
        self.assertEqual(wb["Sources"].max_row, n_sources + 1)
        self.assertEqual(wb["Summary"]["B10"].value, "YES - not a real company")

    # 5: record saved
    def test_record_saved(self):
        db = sqlite3.connect(self.s2.path("general.database"))
        self.addCleanup(db.close)
        n = db.execute("SELECT COUNT(*) FROM analyses").fetchone()[0]
        self.assertEqual(n, 2)
        self.assertGreater(db.execute("SELECT COUNT(*) FROM metrics WHERE analysis_id=1").fetchone()[0], 50)
        self.assertGreater(db.execute("SELECT COUNT(*) FROM sources WHERE analysis_id=1").fetchone()[0], 10)
        self.assertGreater(db.execute("SELECT COUNT(*) FROM conditions WHERE analysis_id=1").fetchone()[0], 3)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM research_notes").fetchone()[0], 2)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM portfolio_imports").fetchone()[0], 1)

    # 6: subsequent analysis compared with the previous one
    def test_comparison(self):
        c = self.changed["comparison"]
        self.assertEqual(c["previous_id"], self.r1["id"])
        self.assertEqual(c["new_periods"], ["Q2FY27"])
        status = {x["id"]: x["status"] for x in c["conditions"]}
        self.assertEqual(status["pat_decline"], "TRIGGERED")     # Q2 PAT 50 vs 56
        self.assertEqual(status["pledge_rise"], "TRIGGERED")     # 4.5% -> 12%
        self.assertEqual(status["support_break"], "TRIGGERED")
        self.assertEqual(status["rev_growth_break"], "holds")    # +3.75% vs threshold 3.19%
        self.assertEqual(status["audit_modified"], "holds")
        moved = {x["metric"] for x in c["metric_changes"]}
        self.assertTrue({"q.pat_yoy", "own.promoter_pledged_pct", "val.price"} <= moved)
        self.assertEqual({a["category"] for a in c["new_announcements"]}, {"financial_results", "kmp_change", "pledge"})
        new_codes = {f["code"] for f in c["findings_new"]}
        self.assertIn("disclosure_kmp_change", new_codes)
        self.assertIn("earnings_decline", new_codes)
        self.assertEqual(self.r2["news"]["new_count"], 3)  # earlier items already seen

    def test_stage2_record_is_clean_including_provenance(self):
        self.assertEqual(self.r2["data_issues"], [])
        for s in self.r2["sources"]:
            if s.get("published"):
                self.assertLessEqual(s["published"][:10], s["retrieved_at"][:10], s["title"])

    def test_no_refresh_comparison_matches(self):
        out = self.agent2.what_changed("SAMPLEMFG", refresh=False)
        self.assertEqual(out["comparison"]["new_periods"], ["Q2FY27"])

    def test_portfolio_fit(self):
        fit = self.r1["portfolio_fit"]
        self.assertTrue(fit["already_held"])
        self.assertEqual(fit["same_sector_holdings"], ["SAMPLEPEERA"])
        self.assertTrue(fit["correlations"])

    def test_peers_and_valuation(self):
        v = self.r1["valuation"]
        self.assertEqual([p["symbol"] for p in v["peers"]], ["SAMPLEPEERA", "SAMPLEPEERB", "SAMPLEPEERC"])
        self.assertEqual(sorted(v["implied"]), ["dcf", "peer_ev_ebitda", "peer_pe"])
        names = {a["name"] for a in self.r1["assumptions"]}
        self.assertIn("risk_free_rate", names)

    def test_ambiguous_query_refused(self):
        with self.assertRaises(ResolutionError) as cm:
            self.agent1.resolve("sample peer")
        self.assertEqual(cm.exception.resolution.status, "ambiguous")

    def test_record_is_json_serialisable(self):
        json.dumps(self.r2, default=str)


class NoPortfolio(unittest.TestCase):
    def test_company_analysis_without_portfolio_or_assumptions(self):
        ws = Workspace()
        try:
            s = ws.settings(1, {"portfolio": {"holdings_files": []}})
            s.data["valuation"].pop("risk_free_rate")
            rec = ResearchAgent(s, as_of=Workspace.as_of(1)).analyse_company("SAMPLEPEERB", write=False)
            self.assertIsNone(rec["portfolio_fit"])
            self.assertTrue(any(m == "dcf" for m, _ in rec["valuation"]["skipped"]))
            codes = {i["code"] for i in rec["data_issues"]}
            self.assertIn("no_announcements", codes)
            self.assertIn("missing", codes)  # no quarterly results / shareholding
            self.assertEqual(rec["summary"]["fundamental_checks_passed"][-1], "4")  # only 3 FY: 3y CAGR check not possible
            self.assertNotIn("revenue_cagr_3y", rec["growth"])
        finally:
            ws.cleanup()


if __name__ == "__main__":
    unittest.main()
